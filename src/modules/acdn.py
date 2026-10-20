#!/usr/bin/env python

import datetime
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from pywebio.output import put_button, put_file, put_loading, put_markdown, put_scope, put_text, use_scope
from pywebio.pin import pin, put_file_upload

from utils import display_random_pet


def _norm_date(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


class Acdn:
    def __init__(self):
        display_random_pet()

        put_markdown("# ACDN 结算明细")

        put_file_upload(
            name="detail_file",
            label="上传结算明细文件",
            accept=".xlsx",
            placeholder="如：XX-X月结算明细.xlsx",
        )
        put_file_upload(
            name="acdn_file",
            label="上传 ACDN 账号级别文件",
            accept=".xlsx",
            placeholder="如：ACDN-账号级别_XXXX.xlsx",
        )
        put_button(
            label="开始匹配文件",
            onclick=self.match_file,
        )

        put_markdown("----")
        put_scope("output")

    @use_scope(name="output", clear=True)
    def match_file(self):
        with use_scope(name="output", clear=True):
            try:
                with put_loading():
                    put_text("开始吭哧吭哧生成结果......")

                    detail: Any = pin.detail_file
                    acdn: Any = pin.acdn_file
                    if detail is None or acdn is None:
                        put_text("请先上传两个文件")
                        return

                    df = pd.read_excel(BytesIO(acdn["content"]), header=1)
                    date_col = df.columns[0]
                    biz_col = df.columns[3]
                    v_col = df.columns[21]
                    x_col = df.columns[23]
                    df["key_date"] = pd.to_datetime(df[date_col], errors="coerce").dt.strftime("%Y-%m-%d")
                    df["key_biz"] = df[biz_col].astype(str).str.strip()
                    df = df.dropna(subset=["key_date"])
                    df = df.rename(columns={v_col: "val_v", x_col: "val_x"})

                    lookup = {(str(r["key_date"]), str(r["key_biz"])): (r["val_v"], r["val_x"]) for r in df.to_dict("records")}

                    wb = load_workbook(BytesIO(detail["content"]))
                    ws = wb.active
                    if ws is None:
                        return

                    ws.cell(row=1, column=12, value=v_col)
                    ws.cell(row=1, column=13, value=x_col)

                    matched = 0
                    unmatched = 0

                    for row in range(2, ws.max_row + 1):
                        biz = str(ws.cell(row=row, column=4).value or "").strip()
                        date_cell = _norm_date(ws.cell(row=row, column=1).value)
                        if not biz or date_cell.startswith("合计"):
                            continue
                        key = (date_cell, biz)
                        if key in lookup:
                            v, x = lookup[key]
                            ws.cell(row=row, column=12, value=None if pd.isna(v) else v)
                            ws.cell(row=row, column=13, value=None if pd.isna(x) else x)
                            matched += 1
                        else:
                            unmatched += 1

                    output_buffer = BytesIO()
                    wb.save(output_buffer)
                    output_buffer.seek(0)
                    content = output_buffer.getvalue()

                put_text(f"匹配完成：共 {matched + unmatched} 行，成功 {matched} 行，未匹配 {unmatched} 行。")
                put_file(
                    name="output.xlsx",
                    content=content,
                    label=">> 点击下载生成后的文件 <<",
                )

            except KeyError as e:
                put_text(f"文件缺少必要列: {e}")
            except Exception as e:
                put_text(f"运行时错误：\n{e}")


if __name__ == "__main__":
    Acdn()
